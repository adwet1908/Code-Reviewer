"""Validate uploads, unpack zips safely, and prepare text for the model.

Uploaded code is never executed. We only decode text and hand it to the LLM.
"""

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from fastapi import UploadFile

from app.config import (
    CODE_EXTENSIONS,
    CODE_MAX_BYTES,
    CONTEXT_EXTENSIONS,
    CONTEXT_MAX_BYTES,
    JUNK_DIR_NAMES,
    MAX_CONTEXT_CHARS,
    MAX_REVIEW_FILES,
    MAX_UPLOADS,
    PROJECT_ZIP_MAX_BYTES,
    PROJECT_ZIP_MAX_UNCOMPRESSED_BYTES,
    ZIP_MAX_BYTES,
    ZIP_MAX_MEMBERS,
    ZIP_MAX_UNCOMPRESSED_BYTES,
)


class IngestError(Exception):
    """A user-facing 4xx. The route turns this into an HTTP error."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


@dataclass
class TextDocument:
    name: str
    text: str
    kind: str  # "code" or "context"


@dataclass
class PreparedBatch:
    code_files: list[TextDocument]
    context_docs: list[TextDocument]


def number_lines(text: str) -> str:
    """Prefix each line so the model can cite a location."""
    lines = text.splitlines()
    width = max(len(str(len(lines))), 1)
    return "\n".join(f"{index:>{width}} | {line}" for index, line in enumerate(lines, start=1))


def format_context(docs: list[TextDocument]) -> str:
    """Join context docs into one block. Truncate so the prompt stays bounded."""
    if not docs:
        return ""

    chunks: list[str] = []
    used = 0
    for doc in docs:
        header = f"### {doc.name}\n"
        room = MAX_CONTEXT_CHARS - used - len(header)
        if room <= 0:
            chunks.append("[further context omitted: size limit]")
            break
        body = doc.text
        if len(body) > room:
            body = body[:room] + "\n[truncated]"
        chunks.append(header + body)
        used += len(chunks[-1]) + 2
    return "\n\n".join(chunks)


def ingest_uploads(files: list[UploadFile]) -> PreparedBatch:
    if not files:
        raise IngestError("Upload at least one file.")
    if len(files) > MAX_UPLOADS:
        raise IngestError(f"Too many files. Upload at most {MAX_UPLOADS} at a time.")

    code_files: list[TextDocument] = []
    context_docs: list[TextDocument] = []
    for upload in files:
        # Browsers sometimes send a path. Keep only the file name.
        raw_name = upload.filename or "upload"
        name = Path(raw_name).name
        data = upload.file.read()
        for doc in ingest_bytes(name, data):
            if doc.kind == "code":
                code_files.append(doc)
            else:
                context_docs.append(doc)

    if len(code_files) > MAX_REVIEW_FILES:
        raise IngestError(
            f"Too many source files ({len(code_files)}). "
            f"The limit is {MAX_REVIEW_FILES} per review."
        )
    if not code_files:
        raise IngestError(
            "No source files to review. Upload code, or a zip that contains source files."
        )
    return PreparedBatch(code_files=code_files, context_docs=context_docs)


def ingest_project_zip_bytes(filename: str, data: bytes) -> list[TextDocument]:
    """Larger zip limits for repo ingestion (Phase 2)."""
    return _read_zip(
        filename,
        data,
        zip_max_bytes=PROJECT_ZIP_MAX_BYTES,
        zip_max_uncompressed=PROJECT_ZIP_MAX_UNCOMPRESSED_BYTES,
    )


def ingest_bytes(filename: str, data: bytes) -> list[TextDocument]:
    """Turn one upload into code and/or context documents."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".zip":
        return _read_zip(filename, data)
    if suffix in CODE_EXTENSIONS:
        _check_size(filename, data, CODE_MAX_BYTES, "Source file")
        text = _decode_text(filename, data)
    elif suffix in CONTEXT_EXTENSIONS:
        _check_size(filename, data, CONTEXT_MAX_BYTES, "Context document")
        text = _extract_context(filename, data)
    else:
        raise IngestError(
            f"Unsupported file type '{suffix or '(none)'}' for {filename}. "
            "Upload source code, a .zip archive, or a .md, .txt, .pdf, or .docx document."
        )
    if not text.strip():
        raise IngestError(f"{filename} is empty.")
    kind = "code" if suffix in CODE_EXTENSIONS else "context"
    return [TextDocument(filename, text, kind)]


def _check_size(filename: str, data: bytes, limit: int, label: str) -> None:
    if len(data) > limit:
        raise IngestError(
            f"{label} {filename} is {len(data)} bytes. The limit is {limit} bytes."
        )


def _decode_text(filename: str, data: bytes) -> str:
    # A null byte means this is not the text file the extension claims.
    if b"\x00" in data:
        raise IngestError(f"{filename} is not valid UTF-8 text.")
    try:
        # utf-8-sig drops a BOM so it does not show up as part of line 1.
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise IngestError(f"{filename} is not valid UTF-8 text.") from None
    return text


def _extract_context(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        return _decode_text(filename, data)
    if suffix == ".pdf":
        return _extract_pdf(filename, data)
    if suffix == ".docx":
        return _extract_docx(filename, data)
    raise IngestError(f"Unsupported context document: {filename}")


def _extract_pdf(filename: str, data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except PdfReadError:
        raise IngestError(f"Could not read PDF {filename}.") from None
    text = "\n\n".join(pages).strip()
    if not text:
        raise IngestError(f"No text found in PDF {filename}. Scanned pages are not supported.")
    return text


def _extract_docx(filename: str, data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(io.BytesIO(data))
    except PackageNotFoundError:
        raise IngestError(f"Could not read DOCX {filename}.") from None
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    text = "\n".join(parts).strip()
    if not text:
        raise IngestError(f"No text found in DOCX {filename}.")
    return text


def _read_zip(
    filename: str,
    data: bytes,
    zip_max_bytes: int = ZIP_MAX_BYTES,
    zip_max_uncompressed: int = ZIP_MAX_UNCOMPRESSED_BYTES,
) -> list[TextDocument]:
    if len(data) > zip_max_bytes:
        raise IngestError(
            f"Zip {filename} is {len(data)} bytes. The limit is {zip_max_bytes} bytes."
        )
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise IngestError(f"{filename} is not a valid zip archive.") from None

    documents: list[TextDocument] = []
    uncompressed_total = 0
    with archive:
        members = archive.infolist()
        if len(members) > ZIP_MAX_MEMBERS:
            raise IngestError(
                f"Zip {filename} has {len(members)} entries. The limit is {ZIP_MAX_MEMBERS}."
            )
        for info in members:
            if info.is_dir():
                continue
            # Bit 0 set means the member is encrypted. We do not ask for passwords.
            if info.flag_bits & 0x1:
                raise IngestError(f"Encrypted zip entries are not supported ({info.filename}).")
            safe_name = _safe_zip_name(info.filename)
            if safe_name is None:
                raise IngestError(f"Zip entry '{info.filename}' has an unsafe path.")
            if _is_junk_path(safe_name):
                continue

            kind = _zip_member_kind(safe_name)
            if kind is None:
                continue

            limit = CODE_MAX_BYTES if kind == "code" else CONTEXT_MAX_BYTES
            label = "Source file" if kind == "code" else "Context document"
            # file_size is only a hint. A zip bomb can lie, so the read below
            # stops after the per-file limit (or whatever uncompressed budget remains).
            if info.file_size > limit:
                raise IngestError(
                    f"{label} {safe_name} is {info.file_size} bytes. The limit is {limit} bytes."
                )
            remaining = zip_max_uncompressed - uncompressed_total
            if remaining <= 0:
                raise IngestError(
                    f"Zip {filename} expands past {zip_max_uncompressed} bytes."
                )
            try:
                payload = _read_limited(archive, info, min(limit, remaining))
            except RuntimeError:
                raise IngestError(f"Could not read zip entry {safe_name}.") from None
            if len(payload) > min(limit, remaining):
                if remaining < limit:
                    raise IngestError(
                        f"Zip {filename} expands past {zip_max_uncompressed} bytes."
                    )
                raise IngestError(
                    f"{label} {safe_name} is over the {limit} byte limit."
                )
            uncompressed_total += len(payload)
            if not payload.strip():
                # Empty __init__.py files are common in repo zips. Skip them.
                continue
            if kind == "code":
                text = _decode_text(safe_name, payload)
            else:
                text = _extract_context(safe_name, payload)
            if not text.strip():
                continue
            documents.append(TextDocument(safe_name, text, kind))

    documents.sort(key=lambda doc: doc.name)
    return documents


def _read_limited(archive: zipfile.ZipFile, info: zipfile.ZipInfo, allowed: int) -> bytes:
    """Read one member, stopping once it passes `allowed` uncompressed bytes."""
    with archive.open(info, "r") as handle:
        return handle.read(allowed + 1)


def _safe_zip_name(name: str) -> str | None:
    """Reject absolute paths and '..' (zip slip). Return a posix relative path."""
    normalized = name.replace("\\", "/")
    if normalized.startswith("/") or normalized.startswith("//"):
        return None
    # "C:/secret.py" is not a relative path inside the archive.
    if len(normalized) >= 2 and normalized[1] == ":":
        return None

    path = PurePosixPath(normalized)
    if path.is_absolute():
        return None
    parts = [part for part in path.parts if part not in ("", ".")]
    if not parts or any(part == ".." for part in parts):
        return None
    return "/".join(parts)


def _is_junk_path(name: str) -> bool:
    return any(part in JUNK_DIR_NAMES for part in PurePosixPath(name).parts)


def _zip_member_kind(name: str) -> str | None:
    """Skip binaries and unknown types inside a zip instead of failing the upload."""
    suffix = PurePosixPath(name).suffix.lower()
    if suffix in CODE_EXTENSIONS:
        return "code"
    if suffix in CONTEXT_EXTENSIONS:
        return "context"
    return None
